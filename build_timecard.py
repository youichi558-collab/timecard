#!/usr/bin/env python3
"""出勤簿ブックの改修スクリプト。

元ブック(出勤簿_2026_original.xlsx)を「その場で」書き換える。作り直すのではなく
元ファイルを開いて必要な箇所だけ直すので、見出し・列幅・フォント・罫線といった
見た目は元のまま変わらない。

直すのは中身だけ:
  - 祝日リストを2025年から2026年に差し替え、VLOOKUPの参照範囲の取りこぼしを修正
  - 勤務時間から実働時間を自動計算し、集計に反映
  - 時刻をコロンなし(930 → 09:30)で入力できるようにする
"""

import datetime
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Border
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SRC = "出勤簿_2026_original.xlsx"
DST = "出勤簿_2026.xlsx"
YEAR = 2026

# 元ブックの列構成(A〜K)は動かさない。右端に2列だけ足す。
C_IN, C_OUT = 4, 6                            # D=出勤, F=退勤
C_OT, C_OT_N, C_OT_E, C_LATE = 7, 8, 9, 10    # G〜J=各残業・遅刻早退
C_REASON = 11                                 # K=理由
C_BREAK, C_WORK = 12, 13                      # L=休憩, M=実働（追加）
# N〜U は非表示の作業列。入力値のシリアル値変換と、曜日区分の判定に使う。
C_CONV = {C_IN: 14, C_OUT: 15, C_BREAK: 16,
          C_OT: 17, C_OT_N: 18, C_OT_E: 19, C_LATE: 20}
C_KIND = 21
HIDDEN = range(14, 22)

INPUT_COLS = [C_IN, C_OUT, C_OT, C_OT_N, C_OT_E, C_LATE, C_BREAK]
COL_LABEL = {C_IN: "出勤", C_OUT: "退勤", C_OT: "通常残業", C_OT_N: "深夜残業",
             C_OT_E: "早朝残業", C_LATE: "遅刻・早退"}

# 時刻入力欄の表示形式。930 という数値を 09:30 と表示する。
# 日付/時刻書式ではなく数値書式にしておくのが要点で、こうしておけば 930 が
# 日付シリアル値と解釈されない。コロン付きで 9:30 と入力した場合は Excel 側が
# 時刻値として取り込み、そのセルの書式を h:mm に置き換える。
TIME_FMT = '00":"00'

HOLIDAYS_2026 = [
    ((1, 1), "元日", None),
    ((1, 12), "成人の日", None),
    ((2, 11), "建国記念の日", None),
    ((2, 23), "天皇誕生日", None),
    ((3, 20), "春分の日", None),
    ((4, 29), "昭和の日", None),
    ((5, 3), "憲法記念日", None),
    ((5, 4), "みどりの日", None),
    ((5, 5), "こどもの日", None),
    ((5, 6), "休日", "祝日法第3条第2項による休日"),
    ((7, 20), "海の日", None),
    ((8, 11), "山の日", None),
    ((9, 21), "敬老の日", None),
    ((9, 22), "国民の休日", "祝日法第3条第3項による休日"),
    ((9, 23), "秋分の日", None),
    ((10, 12), "スポーツの日", None),
    ((11, 3), "文化の日", None),
    ((11, 23), "勤労感謝の日", None),
]
HOLIDAY_NAMES = {name for _, name, _ in HOLIDAYS_2026}

# 元ブックで欄に残っていた飾り文字・見出し語。データではないので消す。
JUNK = {"：", "～", "日", "時間", "出勤日数", "普通出勤時間", "残業", "土曜出勤", "合計"}


def to_serial(ref):
    """時刻入力欄をシリアル値(1日=1)に変換する数式を返す。

    コロンなしで 930 と入れた場合は 9時30分として換算し、コロン付きで 9:30 と
    入れた場合は Excel が時刻値(1未満)として持つのでそのまま使う。
    """
    return (f'IF(NOT(ISNUMBER({ref})),"",'
            f'IF({ref}<1,{ref},(INT({ref}/100)*60+MOD({ref},100))/1440))')


def hhmm(t):
    """datetime.time を、この帳簿の入力形式である HHMM の整数にする。"""
    return t.hour * 100 + t.minute


def copy_style(dst, src):
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


def with_right(border, style):
    """右罫線だけ差し替えた罫線を返す。L・M列を足すぶん、K列の右端を内側の線にする。"""
    right = copy(border.right)
    right.style = style
    return Border(left=copy(border.left), right=right,
                  top=copy(border.top), bottom=copy(border.bottom))


def find_row(ws, col, text, limit=60):
    for r in range(1, limit):
        if ws.cell(r, col).value == text:
            return r
    raise ValueError(f"{ws.title}: 「{text}」の行が見つかりません")


def find_anchor(ws):
    """明細1日目の行を探す。4月だけ開始行が違うため決め打ちにしない。"""
    for r in range(1, 60):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith("=DATE("):
            return r
    raise ValueError(f"{ws.title}: 明細開始行が見つかりません")


def fix_holidays(wb):
    """祝日リストを2026年のものに差し替える。書式は元の行のものを使い回す。"""
    ws = wb["祝日リスト"]
    styles = []
    for c in range(1, 5):
        src = ws.cell(2, c)
        styles.append((copy(src.font), copy(src.border), copy(src.fill),
                       copy(src.alignment), src.number_format))

    for r in range(2, ws.max_row + 1):          # 旧データ(2025年分)を消す
        for c in range(1, 5):
            ws.cell(r, c).value = None

    for i, ((m, d), name, memo) in enumerate(HOLIDAYS_2026, start=2):
        vals = [datetime.date(YEAR, m, d),
                f'=IF(A{i}="","",TEXT(A{i},"（aaa）"))', name, memo]
        for c, (v, st) in enumerate(zip(vals, styles), start=1):
            cell = ws.cell(i, c, v)
            (cell.font, cell.border, cell.fill,
             cell.alignment, cell.number_format) = st
    return len(HOLIDAYS_2026)


def fix_month(ws, name):
    anchor = find_anchor(ws)
    last = anchor + 30
    salvaged = []

    # --- 氏名。元ブックでは月ごとに手入力で、未記入や㊞のままの月があった ---
    if ws.cell(2, C_REASON).value in (None, "", "㊞"):
        ws.cell(2, C_REASON, name)

    # --- 追加する2列の見出し。書式は隣の列から借りて元の体裁に合わせる ---
    head = anchor - 1
    copy_style(ws.cell(head, C_BREAK), ws.cell(head, C_LATE))
    copy_style(ws.cell(head, C_WORK), ws.cell(head, C_REASON))
    ws.cell(head, C_REASON).border = with_right(ws.cell(head, C_REASON).border, "thin")
    ws.cell(head, C_REASON, "理由")   # 4月だけ「休憩時間」になっていた
    ws.cell(head, C_BREAK, "休憩")
    ws.cell(head, C_WORK, "実働")

    for r in range(anchor, last + 1):
        # 曜日欄。4月の1日目だけ数式でなく文字が直接入っていた
        ws.cell(r, 3, f'=IF(A{r}="","","("&TEXT(A{r},"aaa")&"）")')

        # 入力欄。時刻値は HHMM の数値に直し、「：」などの飾り文字は消す
        for c in (C_IN, C_OUT, C_OT, C_OT_N, C_OT_E, C_LATE):
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(v, datetime.time):
                cell.value = hhmm(v)
            elif v is not None:
                if isinstance(v, str) and v.strip() and v.strip() not in JUNK:
                    salvaged.append((r - anchor + 1, cell.coordinate,
                                     f"{COL_LABEL[c]}「{v.strip()}」"))
                cell.value = None
            cell.number_format = TIME_FMT

        # 理由欄。祝日名は自動表示に戻し、手書きのメモはそのまま残す
        k = ws.cell(r, C_REASON)
        keep = (isinstance(k.value, str) and k.value.strip()
                and not k.value.startswith("=")
                and k.value.strip() not in HOLIDAY_NAMES)
        row_bad = [s for d, _, s in salvaged if d == r - anchor + 1]
        if row_bad:
            # 時刻として読めなかった入力は消したままにせず、理由欄に退避する
            flag = "要確認: " + "・".join(row_bad)
            k.value = f"{k.value} / {flag}" if keep else flag
        elif not keep:
            k.value = (f'=IF($A{r}="","",'
                       f'IFERROR(VLOOKUP($A{r},祝日リスト!$A$2:$C$40,3,0),""))')
        k.border = with_right(k.border, "thin")

        # 追加2列。書式は隣の列から借りる
        copy_style(ws.cell(r, C_BREAK), ws.cell(r, C_LATE))
        copy_style(ws.cell(r, C_WORK), ws.cell(r, C_REASON))
        ws.cell(r, C_BREAK).number_format = TIME_FMT
        ws.cell(r, C_WORK).number_format = "[h]:mm"
        ws.cell(r, C_WORK).value = (
            f'=IF(OR($N{r}="",$O{r}=""),"",'
            f'MAX(0,MOD($O{r}-$N{r},1)-IF($P{r}<>"",$P{r},0)))')

        # 作業列。元ブックに残っていた余計な数式もここで上書きされて消える
        for src, dst in C_CONV.items():
            ws.cell(r, dst, "=" + to_serial(f"${get_column_letter(src)}{r}"))
        ws.cell(r, C_KIND,
                f'=IF($A{r}="","",IF($K{r}<>"","祝",'
                f'IF(WEEKDAY($A{r},2)=6,"土",IF(WEEKDAY($A{r},2)=7,"日","平"))))')
        for c in range(C_KIND + 1, 30):         # 作業列より右の残骸を掃除
            ws.cell(r, c).value = None

    # --- 集計。行の位置と見出しは元のまま、式だけ実績ベースに直す ---
    s = find_row(ws, 3, "出勤日数")
    work = f"$M${anchor}:$M${last}"
    kind = f"$U${anchor}:$U${last}"
    for i in range(1, 4):                        # 元は空欄だった行の書式をそろえる
        copy_style(ws.cell(s + i, 4), ws.cell(s, 4))
    ws.cell(s, 4, f"=COUNT($N${anchor}:$N${last})").number_format = "0"
    ws.cell(s + 1, 4, f'=SUM({work})-SUMIF({kind},"土",{work})')
    ws.cell(s + 2, 4,
            f"=SUM($Q${anchor}:$Q${last})+SUM($R${anchor}:$R${last})"
            f"+SUM($S${anchor}:$S${last})")
    ws.cell(s + 3, 4, f'=SUMIF({kind},"土",{work})')
    for i in range(1, 4):
        ws.cell(s + i, 4).number_format = "[h]:mm"
    ws.cell(s + 3, 10, f"=SUM({work})").number_format = "[h]:mm"

    # --- 入力規則。コロンなしとコロン付きの両方を通す ---
    dv = DataValidation(
        type="custom", allow_blank=True, showErrorMessage=True,
        formula1=f'=OR(AND(D{anchor}>=0,D{anchor}<1),'
                 f'AND(D{anchor}=INT(D{anchor}),D{anchor}<=2359,'
                 f'MOD(D{anchor},100)<60))')
    dv.errorTitle = "時刻の入力"
    dv.error = ("コロンなしで 930（＝9時30分）のように入力してください。"
                "9:30 のようにコロン付きでも入力できます。")
    dv.promptTitle = "時刻の入力"
    dv.prompt = "930 と入力すれば 09:30 になります"
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    for c in INPUT_COLS:
        col = get_column_letter(c)
        dv.add(f"{col}{anchor}:{col}{last}")

    # --- 表示・印刷。元の設定に追加2列ぶんだけ足す ---
    for c in HIDDEN:
        ws.column_dimensions[get_column_letter(c)].hidden = True
    ws.column_dimensions["L"].width = 9.0
    ws.column_dimensions["L"].hidden = False
    ws.column_dimensions["M"].width = 10.0
    ws.column_dimensions["M"].hidden = False
    ws.print_area = f"A1:M{s + 3}"
    return salvaged


def main():
    wb = load_workbook(SRC)

    name = ""
    for m in range(1, 13):
        v = wb[f"{m}月"].cell(2, C_REASON).value
        if isinstance(v, str) and v.strip() and v.strip() != "㊞":
            name = v
            break

    n = fix_holidays(wb)
    print(f"祝日リスト: {YEAR}年 {n}件に差し替え")

    for m in range(1, 13):
        for day, addr, v in fix_month(wb[f"{m}月"], name):
            print(f"  要確認 {m}月{day}日 {addr} = {v!r}（時刻として読めないため削除）")

    wb.save(DST)
    print(f"saved {DST}: {wb.sheetnames}")


if __name__ == "__main__":
    main()
